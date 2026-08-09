from __future__ import annotations

import csv
import io
from copy import copy
from collections import defaultdict
from datetime import timedelta, timezone as datetime_timezone

from django.contrib.auth.decorators import login_required
from django.db.models import Count, Max, Sum
from django.http import HttpResponse, JsonResponse
from django.shortcuts import render
from django.utils import timezone

from .core_views import filtered_qsos
from .i18n import catalog, locale_from_request
from .models import QSO
from .utils import histogram_bucket_seconds, histogram_label, range_bounds, qso_duration


def _talkgroup_rows(qsos):
    rows = qsos.values(
        "talkgroup_id",
        "talkgroup__name",
        "talkgroup__country",
        "talkgroup__full_country_name",
        "talkgroup__continent",
    ).annotate(
        qso_count=Count("session_id"),
        duration_ms=Sum("duration_ms"),
        unique_sources=Count("source_call", distinct=True),
        last_seen=Max("start_at"),
    ).order_by("-qso_count", "talkgroup__name")
    return [
        {
            "talkgroup_id": row["talkgroup_id"],
            "name": row["talkgroup__name"] or str(row["talkgroup_id"]),
            "country_code": row["talkgroup__country"] or "XX",
            "country": row["talkgroup__full_country_name"] or "Other",
            "continent": row["talkgroup__continent"] or "Other",
            "qso_count": row["qso_count"],
            "duration_seconds": (row["duration_ms"] or 0) / 1000,
            "unique_sources": row["unique_sources"],
            "last_seen": row["last_seen"].isoformat() if row["last_seen"] else None,
        }
        for row in rows
    ]


def _callsign_rows(qsos):
    # Missing source callsigns are valid raw/QSO data, but they are not a
    # meaningful callsign entry for the report. Do not turn them into a
    # visible synthetic "—" callsign.
    valid_qsos = qsos.exclude(source_call__isnull=True).exclude(source_call__exact="").exclude(source_call__iexact="unknown").exclude(source_call__exact="—")
    rows = list(valid_qsos.values("source_call").annotate(
        qso_count=Count("session_id"),
        duration_ms=Sum("duration_ms"),
        unique_talkgroups=Count("talkgroup_id", distinct=True),
        source_name=Max("source_name"),
    ).order_by("-qso_count", "source_call"))
    countries = defaultdict(set)
    for callsign, country in valid_qsos.values_list("source_call", "talkgroup__full_country_name").distinct().iterator(chunk_size=2000):
        if country:
            countries[callsign].add(country)
    return [
        {
            "callsign": row["source_call"],
            "source_name": row["source_name"] or "",
            "qso_count": row["qso_count"],
            "duration_seconds": (row["duration_ms"] or 0) / 1000,
            "unique_talkgroups": row["unique_talkgroups"],
            "countries": ", ".join(sorted(countries[row["source_call"]])),
        }
        for row in rows
    ]


def _report_rows(request):
    if not request.GET.get("timeRange"):
        request.GET = request.GET.copy()
        request.GET["timeRange"] = "1M"
    qsos = filtered_qsos(request)
    metadata = catalog(locale_from_request(request)).get("metadata", {})
    continent_labels = metadata.get("continents", {})
    country_labels = metadata.get("countries", {})
    start, end = range_bounds(request.GET.get("timeRange"))
    finish = end or timezone.now()
    bucket = histogram_bucket_seconds(request.GET.get("timeRange"))
    daily = defaultdict(lambda: {"qso_count": 0, "duration_seconds": 0})
    hourly = defaultdict(int)
    weekday = defaultdict(int)
    bucket_count = max(1, min(120, int((finish - start).total_seconds() / max(bucket, 1)) + 1))
    histogram_counts = [0] * bucket_count
    active_talkgroups = [set() for _ in range(bucket_count)]
    active_sources = [set() for _ in range(bucket_count)]
    for started_at, talkgroup_id, source_call, duration_ms in qsos.order_by("start_at").values_list("start_at", "talkgroup_id", "source_call", "duration_ms").iterator(chunk_size=5000):
        duration = (duration_ms or 0) / 1000
        started_at_utc = started_at.astimezone(datetime_timezone.utc)
        key = started_at_utc.date().isoformat()
        daily[key]["qso_count"] += 1
        daily[key]["duration_seconds"] += duration
        hourly[started_at_utc.hour] += 1
        weekday[started_at_utc.isoweekday()] += 1
        index = int((started_at - start).total_seconds() / max(bucket, 1))
        if 0 <= index < bucket_count:
            histogram_counts[index] += 1
            active_talkgroups[index].add(talkgroup_id)
            active_sources[index].add(source_call)
    summary = qsos.aggregate(qso_count=Count("session_id"), duration_ms=Sum("duration_ms"), unique_sources=Count("source_call", distinct=True), unique_talkgroups=Count("talkgroup_id", distinct=True))
    current_count = summary["qso_count"] or 0
    talkgroup_rows = _talkgroup_rows(qsos)
    callsign_rows = _callsign_rows(qsos)
    callsigns_by_duration = sorted(
        callsign_rows,
        key=lambda row: (-row["duration_seconds"], row["callsign"]),
    )
    selector_query = request.GET.copy()
    selector_query.pop("callsign", None)
    selector_query.pop("talkgroup", None)
    selector_request = copy(request)
    selector_request.GET = selector_query
    selector_talkgroups = _talkgroup_rows(filtered_qsos(selector_request))
    selector_continents = sorted(
        {(value, continent_labels.get(value, value)) for value in {row["continent"] for row in selector_talkgroups}},
        key=lambda row: row[1].casefold(),
    )
    selector_countries = sorted(
        {(code, country_labels.get(code, name)) for code, name in {(row["country_code"], row["country"]) for row in selector_talkgroups}},
        key=lambda row: row[1].casefold(),
    )
    histogram = []
    for index, qso_count in enumerate(histogram_counts):
        cursor = start + timedelta(seconds=index * bucket)
        histogram.append({"bucket": cursor.isoformat(), "label": histogram_label(cursor, bucket), "qso_count": qso_count})
    previous_start = start - (finish - start)
    previous_end = start
    previous_count = QSO.objects.filter(start_at__gte=previous_start, start_at__lt=previous_end).count()
    trend = None if previous_count == 0 else round((current_count - previous_count) * 100 / previous_count, 1)
    growth = []
    previous = {
        row["talkgroup_id"]: row["qso_count"]
        for row in QSO.objects.filter(start_at__gte=previous_start, start_at__lt=previous_end)
        .values("talkgroup_id")
        .annotate(qso_count=Count("session_id"))
    }
    for row in talkgroup_rows:
        old = previous.get(row["talkgroup_id"], 0)
        row_copy = {"talkgroup_id": row["talkgroup_id"], "name": row["name"], "current_qso_count": row["qso_count"], "previous_qso_count": old, "growth_percent": None if old == 0 else round((row["qso_count"] - old) * 100 / old, 1)}
        growth.append(row_copy)
    growth.sort(key=lambda row: (row["growth_percent"] is not None, row["growth_percent"] or 0, row["current_qso_count"]), reverse=True)
    concurrency = [
        {"bucket": (cursor := start + timedelta(seconds=index * bucket)).isoformat(), "label": histogram_label(cursor, bucket), "active_talkgroups": len(active_talkgroups[index]), "active_sources": len(active_sources[index])}
        for index in range(bucket_count)
    ]
    return {"summary": {"qso_count": current_count, "duration_seconds": (summary["duration_ms"] or 0) / 1000, "unique_sources": summary["unique_sources"] or 0, "unique_talkgroups": summary["unique_talkgroups"] or 0}, "daily": [{"day": key, **value} for key, value in sorted(daily.items())][:50], "histogram": histogram, "histogram_max_qsos": max(histogram_counts, default=1) or 1, "talkgroups": talkgroup_rows[:50], "selector_talkgroups": selector_talkgroups, "selector_continents": selector_continents, "selector_countries": selector_countries, "callsigns": callsign_rows[:50], "callsigns_by_duration": callsigns_by_duration[:50], "hourly_activity": [{"hour": key, "qso_count": hourly[key]} for key in range(24)], "weekday_activity": [{"weekday": key, "qso_count": weekday[key]} for key in range(1, 8)], "peak_periods": sorted(histogram, key=lambda row: row["qso_count"], reverse=True)[:5], "traffic_trend": {"current_qsos": current_count, "previous_qsos": previous_count, "qso_change_percent": trend}, "talkgroup_growth": growth[:50], "concurrent_activity": concurrency, "concurrent_max_talkgroups": max((row["active_talkgroups"] for row in concurrency), default=1) or 1, "concurrent_max_sources": max((row["active_sources"] for row in concurrency), default=1) or 1, "bucket_seconds": bucket}


@login_required
def reports(request):
    data = _report_rows(request)
    ranges = [("24h", "Last 24 hours"), ("today", "Today"), ("yesterday", "Yesterday"), ("1w", "Last 7 days"), ("lastWeek", "Last week"), ("2w", "Last 14 days"), ("1M", "Last 30 days"), ("lastMonth", "Last month"), ("2M", "Last 2 months"), ("3M", "Last 3 months")]
    return render(request, "dashboard/reports.html", {"report": data, "ranges": ranges, "time_range": request.GET.get("timeRange", "1M"), "selected_talkgroups": request.GET.getlist("talkgroup"), "records_retrieved": data["summary"]["qso_count"], "query_seconds": "—"})


@login_required
def report_api(request):
    return JsonResponse(_report_rows(request))


def _report_for_export(request):
    return _report_rows(request)


@login_required
def report_csv(request):
    data = _report_for_export(request)
    output = io.StringIO(); writer = csv.writer(output)
    writer.writerow(["Talkgroup", "QSOs", "Talk time", "Unique sources", "Last heard"])
    for row in data["talkgroups"]: writer.writerow([f"{row['name']} ({row['talkgroup_id']})", row["qso_count"], qso_duration(row["duration_seconds"]), row["unique_sources"], row["last_seen"]])
    writer.writerow([]); writer.writerow(["Callsign", "Name", "Countries", "QSOs", "Talk time", "Talkgroups"])
    for row in data["callsigns"]: writer.writerow([row["callsign"], row["source_name"], row["countries"], row["qso_count"], qso_duration(row["duration_seconds"]), row["unique_talkgroups"]])
    response = HttpResponse(output.getvalue(), content_type="text/csv; charset=utf-8"); response["Content-Disposition"] = "attachment; filename=brandmeister-report.csv"; return response


@login_required
def report_excel(request):
    from openpyxl import Workbook
    data = _report_for_export(request); workbook = Workbook(); sheet = workbook.active; sheet.title = "Talkgroups"
    sheet.append(["Talkgroup", "QSOs", "Talk time", "Unique sources", "Last heard"])
    for row in data["talkgroups"]: sheet.append([f"{row['name']} ({row['talkgroup_id']})", row["qso_count"], qso_duration(row["duration_seconds"]), row["unique_sources"], row["last_seen"]])
    calls = workbook.create_sheet("Callsigns"); calls.append(["Callsign", "Name", "Countries", "QSOs", "Talk time", "Talkgroups"])
    for row in data["callsigns"]: calls.append([row["callsign"], row["source_name"], row["countries"], row["qso_count"], qso_duration(row["duration_seconds"]), row["unique_talkgroups"]])
    output = io.BytesIO(); workbook.save(output); response = HttpResponse(output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"); response["Content-Disposition"] = "attachment; filename=brandmeister-report.xlsx"; return response


def _pdf_histogram(rows, width):
    """Create a proportional vector histogram for the selected report period."""
    from reportlab.graphics.shapes import Drawing, Line, Rect, String
    from reportlab.lib.colors import HexColor

    rows = list(rows)
    if not rows:
        return None
    height = 178
    left, bottom, top = 32, 38, 18
    chart_width = max(10, width - left - 8)
    chart_height = height - bottom - top
    maximum = max(int(row.get("qso_count") or 0) for row in rows) or 1
    drawing = Drawing(width, height)
    axis = HexColor("#94a3b8")
    muted = HexColor("#64748b")
    accent = HexColor("#667eea")
    drawing.add(Line(left, bottom, left + chart_width, bottom, strokeColor=axis, strokeWidth=0.7))
    drawing.add(String(left - 4, bottom + chart_height - 2, str(maximum), fontSize=7, fillColor=muted, textAnchor="end"))
    drawing.add(String(left - 4, bottom - 2, "0", fontSize=7, fillColor=muted, textAnchor="end"))
    bar_width = chart_width / max(len(rows), 1)
    label_step = max(1, (len(rows) + 9) // 10)
    for index, row in enumerate(rows):
        count = int(row.get("qso_count") or 0)
        x = left + index * bar_width + 0.5
        bar_height = count / maximum * chart_height if count else 0
        if bar_height:
            drawing.add(Rect(x, bottom, max(1, bar_width - 1), bar_height, fillColor=accent, strokeColor=None, rx=1, ry=1))
            drawing.add(String(x + bar_width / 2, bottom + bar_height + 3, str(count), fontSize=6, fillColor=muted, textAnchor="middle"))
        if index % label_step == 0 or index == len(rows) - 1:
            label = str(row.get("label") or "")
            drawing.add(String(x + bar_width / 2, 22, label, fontSize=6, fillColor=muted, textAnchor="middle"))
    return drawing


def _pdf_horizontal_bars(rows, label_key, value_key, width, formatter=None, limit=12, color="#764ba2"):
    """Create a readable proportional horizontal bar chart."""
    from reportlab.graphics.shapes import Drawing, Rect, String
    from reportlab.lib.colors import HexColor

    rows = list(rows)[:limit]
    if not rows:
        return None
    row_height = 20
    height = row_height * len(rows) + 8
    label_width = min(168, max(112, width * 0.32))
    value_width = 48
    chart_width = max(30, width - label_width - value_width)
    maximum = max(float(row.get(value_key) or 0) for row in rows) or 1
    drawing = Drawing(width, height)
    for index, row in enumerate(rows):
        y = height - (index + 1) * row_height + 5
        label = str(row.get(label_key) or "-")
        if len(label) > 28:
            label = label[:25] + "..."
        amount = float(row.get(value_key) or 0)
        drawing.add(String(0, y + 2, label, fontSize=7, fillColor=HexColor("#334155")))
        if amount:
            drawing.add(Rect(label_width, y, max(1.2, amount / maximum * chart_width), 10, fillColor=HexColor(color), strokeColor=None, rx=2, ry=2))
        display = formatter(amount) if formatter else str(int(amount))
        drawing.add(String(width, y + 2, display, fontSize=7, fillColor=HexColor("#64748b"), textAnchor="end"))
    return drawing


def _pdf_concurrency_chart(rows, width):
    """Create a common-scale two-series chart with values above each bar."""
    from reportlab.graphics.shapes import Drawing, Line, Rect, String
    from reportlab.lib.colors import HexColor

    rows = list(rows)
    if not rows:
        return None
    height = 190
    left, bottom, top = 32, 42, 24
    chart_width = max(10, width - left - 8)
    chart_height = height - bottom - top
    maximum = max(
        max((int(row.get("active_talkgroups") or 0) for row in rows), default=0),
        max((int(row.get("active_sources") or 0) for row in rows), default=0),
    ) or 1
    drawing = Drawing(width, height)
    muted = HexColor("#64748b")
    drawing.add(Line(left, bottom, left + chart_width, bottom, strokeColor=HexColor("#94a3b8"), strokeWidth=0.7))
    drawing.add(String(left - 4, bottom + chart_height - 2, str(maximum), fontSize=7, fillColor=muted, textAnchor="end"))
    drawing.add(String(left - 4, bottom - 2, "0", fontSize=7, fillColor=muted, textAnchor="end"))
    group_width = chart_width / max(len(rows), 1)
    label_step = max(1, (len(rows) + 9) // 10)
    for index, row in enumerate(rows):
        group_x = left + index * group_width
        bar_width = max(1.3, min(10, group_width * 0.32))
        gap = max(1, group_width * 0.04)
        talkgroups = int(row.get("active_talkgroups") or 0)
        sources = int(row.get("active_sources") or 0)
        for offset, amount, color in ((-bar_width - gap / 2, talkgroups, "#667eea"), (gap / 2, sources, "#a855f7")):
            x = group_x + group_width / 2 + offset
            bar_height = amount / maximum * chart_height if amount else 0
            if bar_height:
                drawing.add(Rect(x, bottom, bar_width, bar_height, fillColor=HexColor(color), strokeColor=None, rx=1, ry=1))
                drawing.add(String(x + bar_width / 2, bottom + bar_height + 3, str(amount), fontSize=5.5, fillColor=muted, textAnchor="middle"))
        if index % label_step == 0 or index == len(rows) - 1:
            drawing.add(String(group_x + group_width / 2, 25, str(row.get("label") or ""), fontSize=6, fillColor=muted, textAnchor="middle", angle=45))
    return drawing


@login_required
def report_pdf(request):
    from xml.sax.saxutils import escape

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, TableStyle

    data = _report_for_export(request)
    output = io.BytesIO()
    page_width, page_height = A4
    margin = 14 * mm
    content_width = page_width - 2 * margin
    doc = SimpleDocTemplate(output, pagesize=A4, rightMargin=margin, leftMargin=margin, topMargin=16 * mm, bottomMargin=16 * mm)
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="ReportSubtitle", parent=styles["Normal"], fontSize=9, leading=12, textColor=colors.HexColor("#64748b"), spaceAfter=4))
    styles.add(ParagraphStyle(name="ReportSection", parent=styles["Heading2"], fontSize=13, leading=16, textColor=colors.HexColor("#273449"), spaceBefore=10, spaceAfter=6, keepWithNext=True))
    styles.add(ParagraphStyle(name="ReportCell", parent=styles["Normal"], fontSize=7.5, leading=9, textColor=colors.HexColor("#334155")))
    styles.add(ParagraphStyle(name="ReportHeaderCell", parent=styles["Normal"], fontSize=7.5, leading=9, textColor=colors.white, fontName="Helvetica-Bold"))

    def cell(value, header=False):
        text = escape(str(value if value not in (None, "") else "-"))
        return Paragraph(text, styles["ReportHeaderCell" if header else "ReportCell"])

    def report_table(headers, rows, widths):
        values = [[cell(value, True) for value in headers]]
        values.extend([[cell(value) for value in row] for row in rows])
        if len(values) == 1:
            values.append([cell("No data")] + [cell("") for _ in headers[1:]])
        table = Table(values, colWidths=widths, repeatRows=1, hAlign="LEFT")
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#667eea")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#dbe3ef")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 5),
            ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        return table

    def compact_countries(value):
        countries = [part.strip() for part in str(value or "").split(",") if part.strip()]
        if len(countries) > 4:
            return ", ".join(countries[:3]) + f", +{len(countries) - 3} more"
        return ", ".join(countries) or "-"

    range_labels = {"24h": "Last 24 hours", "today": "Today", "yesterday": "Yesterday", "1w": "Last 7 days", "lastWeek": "Last week", "2w": "Last 14 days", "1M": "Last 30 days", "lastMonth": "Last month", "2M": "Last 2 months", "3M": "Last 3 months"}
    selected_range = range_labels.get(request.GET.get("timeRange"), "Selected period")
    filters = [f"Period: {selected_range}"]
    for key, label in (("continent", "Continent"), ("country", "Country"), ("callsign", "Callsign")):
        value = request.GET.get(key, "").strip()
        if value:
            filters.append(f"{label}: {value}")
    selected_talkgroups = request.GET.getlist("talkgroup")
    if selected_talkgroups:
        filters.append(f"Talkgroups: {', '.join(selected_talkgroups)}")

    summary = data["summary"]
    trend = data["traffic_trend"]["qso_change_percent"]
    trend_text = f"{trend}%" if trend is not None else "N/A"
    story = [
        Paragraph("BrandMeister Lastheard report", styles["Title"]),
        Paragraph(escape(" / ".join(filters)), styles["ReportSubtitle"]),
    ]
    story.append(report_table(
        ["QSOs", "Talk time", "Callsigns", "Talkgroups", "Traffic trend"],
        [[summary["qso_count"], qso_duration(summary["duration_seconds"]), summary["unique_sources"], summary["unique_talkgroups"], f"{trend_text} (current {data['traffic_trend']['current_qsos']} / previous {data['traffic_trend']['previous_qsos']})"]],
        [content_width * ratio for ratio in (0.16, 0.19, 0.16, 0.17, 0.32)],
    ))

    story.append(Paragraph("Daily activity", styles["ReportSection"]))
    daily_chart = _pdf_histogram(data["histogram"], content_width)
    if daily_chart:
        story.append(daily_chart)

    story.append(Paragraph("Talkgroup activity", styles["ReportSection"]))
    talkgroup_chart_rows = [{"label": f"{row['name']} ({row['talkgroup_id']})", "qso_count": row["qso_count"]} for row in data["talkgroups"]]
    talkgroup_chart = _pdf_horizontal_bars(talkgroup_chart_rows, "label", "qso_count", content_width)
    if talkgroup_chart:
        story.append(talkgroup_chart)
    story.append(report_table(
        ["Talkgroup", "QSOs", "Talk time", "Sources", "Last heard"],
        [[f"{row['name']} ({row['talkgroup_id']})", row["qso_count"], qso_duration(row["duration_seconds"]), row["unique_sources"], row["last_seen"] or "-"] for row in data["talkgroups"]],
        [content_width * ratio for ratio in (0.32, 0.11, 0.16, 0.12, 0.29)],
    ))

    story.append(Paragraph("Most active callsigns by QSOs", styles["ReportSection"]))
    qso_callsign_rows = [{"label": f"{row['callsign']}" + (f" - {row['source_name']}" if row["source_name"] else ""), "qso_count": row["qso_count"]} for row in data["callsigns"]]
    qso_callsign_chart = _pdf_horizontal_bars(qso_callsign_rows, "label", "qso_count", content_width)
    if qso_callsign_chart:
        story.append(qso_callsign_chart)

    story.append(Paragraph("Most active callsigns by talk time", styles["ReportSection"]))
    duration_callsign_rows = [{"label": f"{row['callsign']}" + (f" - {row['countries']}" if row["countries"] else ""), "duration_seconds": row["duration_seconds"]} for row in data["callsigns_by_duration"]]
    duration_callsign_chart = _pdf_horizontal_bars(duration_callsign_rows, "label", "duration_seconds", content_width, formatter=lambda value: qso_duration(value))
    if duration_callsign_chart:
        story.append(duration_callsign_chart)
    story.append(report_table(
        ["Callsign", "Name", "Country", "QSOs", "Talk time", "Talkgroups"],
        [[row["callsign"], row["source_name"] or "-", compact_countries(row["countries"]), row["qso_count"], qso_duration(row["duration_seconds"]), row["unique_talkgroups"]] for row in data["callsigns"]],
        [content_width * ratio for ratio in (0.13, 0.20, 0.25, 0.10, 0.17, 0.15)],
    ))

    story.append(Paragraph("QSOs by hour", styles["ReportSection"]))
    hourly_rows = [{"label": f"{row['hour']:02d}:00", "qso_count": row["qso_count"]} for row in data["hourly_activity"]]
    hourly_chart = _pdf_horizontal_bars(hourly_rows, "label", "qso_count", content_width, limit=24)
    if hourly_chart:
        story.append(hourly_chart)

    story.append(Paragraph("QSOs by weekday", styles["ReportSection"]))
    weekday_names = ("Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday")
    weekday_rows = [{"label": weekday_names[int(row["weekday"]) - 1] if 1 <= int(row["weekday"]) <= 7 else str(row["weekday"]), "qso_count": row["qso_count"]} for row in data["weekday_activity"]]
    weekday_chart = _pdf_horizontal_bars(weekday_rows, "label", "qso_count", content_width, limit=7)
    if weekday_chart:
        story.append(weekday_chart)

    story.append(Paragraph("Concurrent active talkgroups and sources", styles["ReportSection"]))
    story.append(Paragraph("Blue: active talkgroups / Purple: active sources", styles["ReportSubtitle"]))
    concurrency_chart = _pdf_concurrency_chart(data["concurrent_activity"], content_width)
    if concurrency_chart:
        story.append(concurrency_chart)

    def draw_page(canvas, document):
        canvas.saveState()
        canvas.setStrokeColor(colors.HexColor("#667eea"))
        canvas.setLineWidth(1.2)
        canvas.line(margin, page_height - 10 * mm, page_width - margin, page_height - 10 * mm)
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor("#64748b"))
        canvas.drawString(margin, 8 * mm, "BrandMeister Lastheard")
        canvas.drawRightString(page_width - margin, 8 * mm, f"Page {canvas.getPageNumber()}")
        canvas.restoreState()

    doc.build(story, onFirstPage=draw_page, onLaterPages=draw_page)
    response = HttpResponse(output.getvalue(), content_type="application/pdf")
    response["Content-Disposition"] = "attachment; filename=brandmeister-report.pdf"
    return response
